from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Depends, UploadFile, File, Query, Header, Response
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Literal
from datetime import datetime, timezone, timedelta
from bson import ObjectId
import logging
import uuid
import io
import bcrypt
import jwt
import requests

# ---------------- DB ----------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGO = "HS256"
TRIAL_DAYS = 14

DEFAULT_PLANS = [
    {"id": "monthly", "name": "Paket Bulanan", "price": 49000, "days": 30, "features": ["Kasir tanpa batas", "Dashboard & Laporan", "1 Outlet"]},
    {"id": "yearly", "name": "Paket Tahunan", "price": 490000, "days": 365, "features": ["Semua fitur Bulanan", "Multi-outlet", "Prioritas dukungan", "Hemat 2 bulan"]},
]

# ---------------- Storage ----------------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "umkmpay"
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# ---------------- Auth helpers ----------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(user_id: str, email: str, role: str) -> str:
    payload = {"sub": user_id, "email": email, "role": role,
               "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


def clean_user(u: dict) -> dict:
    u = dict(u)
    u["id"] = str(u.pop("_id"))
    u.pop("password_hash", None)
    return u


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        if not user.get("is_active", True):
            raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
        return clean_user(user)
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login lagi")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")


def require_roles(*roles):
    async def dep(user: dict = Depends(get_current_user)):
        if user["role"] not in roles:
            raise HTTPException(status_code=403, detail="Akses ditolak")
        return user
    return dep


# ---------------- Models ----------------
class RegisterInput(BaseModel):
    name: str
    business_name: str
    email: EmailStr
    password: str
    phone: Optional[str] = ""


class LoginInput(BaseModel):
    email: EmailStr
    password: str


class ProductInput(BaseModel):
    name: str
    sku: Optional[str] = ""
    category: Optional[str] = "Umum"
    price: float
    cost: float = 0
    stock: int = 0
    low_stock_threshold: int = 5
    image_path: Optional[str] = None
    outlet_id: Optional[str] = None


class SaleItem(BaseModel):
    product_id: str
    name: str
    price: float
    cost: float = 0
    qty: int


class SaleInput(BaseModel):
    items: List[SaleItem]
    payment_method: Literal["cash", "qris"] = "cash"
    tax_rate: float = 0
    discount: float = 0
    customer_id: Optional[str] = None
    is_credit: bool = False
    amount_paid: float = 0
    outlet_id: Optional[str] = None
    note: Optional[str] = ""


class ExpenseInput(BaseModel):
    category: str
    amount: float
    note: Optional[str] = ""
    outlet_id: Optional[str] = None


class CustomerInput(BaseModel):
    name: str
    phone: Optional[str] = ""
    note: Optional[str] = ""


class OutletInput(BaseModel):
    name: str
    address: Optional[str] = ""


class SupplierInput(BaseModel):
    name: str
    phone: Optional[str] = ""
    note: Optional[str] = ""


class PurchaseItem(BaseModel):
    product_id: str
    name: str
    cost: float
    qty: int


class PurchaseInput(BaseModel):
    items: List[PurchaseItem]
    supplier_id: Optional[str] = None
    outlet_id: Optional[str] = None
    note: Optional[str] = ""
    payment_method: Literal["cash", "qris", "credit"] = "cash"


class CashierInput(BaseModel):
    name: str
    email: EmailStr
    password: str
    outlet_id: Optional[str] = None


class SubscribeInput(BaseModel):
    plan_id: str
    proof_path: str


class UmkmProfileInput(BaseModel):
    business_name: Optional[str] = None
    qris_image_path: Optional[str] = None
    logo_image_path: Optional[str] = None
    address: Optional[str] = None
    phone: Optional[str] = None


class PlatformSettingsInput(BaseModel):
    qris_image_path: Optional[str] = None
    business_name: Optional[str] = None
    contact: Optional[str] = None


def now_iso():
    return datetime.now(timezone.utc).isoformat()


async def get_umkm(umkm_id: str):
    return await db.umkms.find_one({"id": umkm_id}, {"_id": 0})


def effective_status(umkm: dict) -> dict:
    end = umkm.get("subscription_end")
    active = False
    if end:
        try:
            active = datetime.fromisoformat(end) > datetime.now(timezone.utc)
        except Exception:
            active = False
    status = "active" if active else "expired"
    if umkm.get("status") == "trial" and active:
        status = "trial"
    return {"status": status, "active": active, "subscription_end": end, "plan": umkm.get("plan")}


# ================= AUTH =================
@api_router.post("/auth/register")
async def register(data: RegisterInput):
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    umkm_id = str(uuid.uuid4())
    trial_end = (datetime.now(timezone.utc) + timedelta(days=TRIAL_DAYS)).isoformat()
    await db.umkms.insert_one({
        "id": umkm_id, "name": data.business_name, "owner_email": email,
        "status": "trial", "plan": "trial", "subscription_end": trial_end,
        "qris_image_path": None, "address": "", "phone": data.phone or "",
        "created_at": now_iso(),
    })
    default_outlet = str(uuid.uuid4())
    await db.outlets.insert_one({"id": default_outlet, "umkm_id": umkm_id, "name": "Outlet Utama", "address": "", "created_at": now_iso()})
    res = await db.users.insert_one({
        "email": email, "password_hash": hash_password(data.password), "name": data.name,
        "role": "umkm_admin", "umkm_id": umkm_id, "outlet_id": default_outlet,
        "is_active": True, "created_at": now_iso(),
    })
    uid = str(res.inserted_id)
    token = create_token(uid, email, "umkm_admin")
    user = await db.users.find_one({"_id": res.inserted_id})
    return {"access_token": token, "user": clean_user(user)}


@api_router.post("/auth/login")
async def login(data: LoginInput):
    email = data.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau password salah")
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    token = create_token(str(user["_id"]), email, user["role"])
    return {"access_token": token, "user": clean_user(user)}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    if user.get("umkm_id"):
        umkm = await get_umkm(user["umkm_id"])
        if umkm:
            user["umkm"] = umkm
            user["subscription"] = effective_status(umkm)
    return user


@api_router.post("/auth/logout")
async def logout(user: dict = Depends(get_current_user)):
    return {"ok": True}


# ================= UPLOAD / FILES =================
MIME = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp", "gif": "image/gif"}


@api_router.post("/upload")
async def upload(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    path = f"{APP_NAME}/uploads/{user['id']}/{uuid.uuid4()}.{ext}"
    data = await file.read()
    ct = file.content_type or MIME.get(ext, "application/octet-stream")
    result = put_object(path, data, ct)
    await db.files.insert_one({"id": str(uuid.uuid4()), "storage_path": result["path"], "original_filename": file.filename,
                               "content_type": ct, "size": result.get("size", len(data)), "is_deleted": False, "created_at": now_iso()})
    return {"path": result["path"]}


@api_router.get("/files/{path:path}")
async def download(path: str, authorization: str = Header(None), auth: str = Query(None)):
    token = None
    if authorization and authorization.startswith("Bearer "):
        token = authorization[7:]
    elif auth:
        token = auth
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    record = await db.files.find_one({"storage_path": path, "is_deleted": False})
    if not record:
        raise HTTPException(status_code=404, detail="File tidak ditemukan")
    content, ct = get_object(path)
    return Response(content=content, media_type=record.get("content_type", ct))


# ================= UMKM PROFILE =================
@api_router.get("/umkm")
async def my_umkm(user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    umkm = await get_umkm(user["umkm_id"])
    if not umkm:
        raise HTTPException(status_code=404, detail="UMKM tidak ditemukan")
    umkm["subscription"] = effective_status(umkm)
    return umkm


@api_router.put("/umkm")
async def update_umkm(data: UmkmProfileInput, user: dict = Depends(require_roles("umkm_admin"))):
    upd = {k: v for k, v in data.model_dump().items() if v is not None}
    if "business_name" in upd:
        upd["name"] = upd.pop("business_name")
    if upd:
        await db.umkms.update_one({"id": user["umkm_id"]}, {"$set": upd})
    return await get_umkm(user["umkm_id"])


# ================= SUBSCRIPTION =================
@api_router.get("/subscription/plans")
async def get_plans():
    settings = await db.platform.find_one({"id": "platform"}, {"_id": 0})
    return {"plans": (settings or {}).get("plans", DEFAULT_PLANS), "qris_image_path": (settings or {}).get("qris_image_path")}


@api_router.get("/subscription")
async def my_subscription(user: dict = Depends(require_roles("umkm_admin"))):
    umkm = await get_umkm(user["umkm_id"])
    payments = await db.subscriptions.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return {"subscription": effective_status(umkm), "payments": payments}


@api_router.post("/subscription/subscribe")
async def subscribe(data: SubscribeInput, user: dict = Depends(require_roles("umkm_admin"))):
    settings = await db.platform.find_one({"id": "platform"}, {"_id": 0}) or {}
    plans = settings.get("plans", DEFAULT_PLANS)
    plan = next((p for p in plans if p["id"] == data.plan_id), None)
    if not plan:
        raise HTTPException(status_code=400, detail="Paket tidak valid")
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "umkm_name": user.get("umkm_id"),
           "plan_id": plan["id"], "plan_name": plan["name"], "amount": plan["price"], "days": plan["days"],
           "proof_path": data.proof_path, "status": "pending", "created_at": now_iso(), "reviewed_at": None}
    umkm = await get_umkm(user["umkm_id"])
    doc["umkm_name"] = umkm["name"]
    await db.subscriptions.insert_one(doc)
    doc.pop("_id", None)
    return doc


# ================= SUPER ADMIN =================
@api_router.get("/admin/stats")
async def admin_stats(user: dict = Depends(require_roles("super_admin"))):
    umkms = await db.umkms.find({}, {"_id": 0}).to_list(1000)
    active = sum(1 for u in umkms if effective_status(u)["active"])
    pending = await db.subscriptions.count_documents({"status": "pending"})
    approved = await db.subscriptions.find({"status": "approved"}, {"_id": 0}).to_list(1000)
    revenue = sum(p.get("amount", 0) for p in approved)
    return {"total_umkm": len(umkms), "active_umkm": active, "pending_payments": pending, "total_revenue": revenue}


@api_router.get("/admin/umkms")
async def admin_umkms(user: dict = Depends(require_roles("super_admin"))):
    umkms = await db.umkms.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    for u in umkms:
        u["subscription"] = effective_status(u)
        u["user_count"] = await db.users.count_documents({"umkm_id": u["id"]})
    return umkms


@api_router.get("/admin/subscriptions")
async def admin_subscriptions(status: Optional[str] = None, user: dict = Depends(require_roles("super_admin"))):
    q = {"status": status} if status else {}
    return await db.subscriptions.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


@api_router.post("/admin/subscriptions/{sub_id}/approve")
async def approve_sub(sub_id: str, user: dict = Depends(require_roles("super_admin"))):
    sub = await db.subscriptions.find_one({"id": sub_id})
    if not sub:
        raise HTTPException(status_code=404, detail="Pembayaran tidak ditemukan")
    umkm = await get_umkm(sub["umkm_id"])
    base = datetime.now(timezone.utc)
    cur_end = umkm.get("subscription_end")
    if cur_end:
        try:
            d = datetime.fromisoformat(cur_end)
            if d > base:
                base = d
        except Exception:
            pass
    new_end = (base + timedelta(days=sub["days"])).isoformat()
    await db.umkms.update_one({"id": sub["umkm_id"]}, {"$set": {"status": "active", "plan": sub["plan_id"], "subscription_end": new_end}})
    await db.subscriptions.update_one({"id": sub_id}, {"$set": {"status": "approved", "reviewed_at": now_iso()}})
    return {"ok": True, "subscription_end": new_end}


@api_router.post("/admin/subscriptions/{sub_id}/reject")
async def reject_sub(sub_id: str, user: dict = Depends(require_roles("super_admin"))):
    r = await db.subscriptions.update_one({"id": sub_id}, {"$set": {"status": "rejected", "reviewed_at": now_iso()}})
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Pembayaran tidak ditemukan")
    return {"ok": True}


@api_router.post("/admin/umkms/{umkm_id}/toggle")
async def toggle_umkm(umkm_id: str, user: dict = Depends(require_roles("super_admin"))):
    umkm = await get_umkm(umkm_id)
    if not umkm:
        raise HTTPException(status_code=404, detail="UMKM tidak ditemukan")
    new = not umkm.get("suspended", False)
    await db.umkms.update_one({"id": umkm_id}, {"$set": {"suspended": new}})
    await db.users.update_many({"umkm_id": umkm_id}, {"$set": {"is_active": not new}})
    return {"ok": True, "suspended": new}


@api_router.get("/admin/settings")
async def get_settings(user: dict = Depends(require_roles("super_admin"))):
    s = await db.platform.find_one({"id": "platform"}, {"_id": 0})
    return s or {"id": "platform", "plans": DEFAULT_PLANS}


@api_router.put("/admin/settings")
async def update_settings(data: PlatformSettingsInput, user: dict = Depends(require_roles("super_admin"))):
    upd = {k: v for k, v in data.model_dump().items() if v is not None}
    await db.platform.update_one({"id": "platform"}, {"$set": upd}, upsert=True)
    return await db.platform.find_one({"id": "platform"}, {"_id": 0})


# ================= OUTLETS =================
@api_router.get("/outlets")
async def list_outlets(user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    return await db.outlets.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).to_list(100)


@api_router.post("/outlets")
async def create_outlet(data: OutletInput, user: dict = Depends(require_roles("umkm_admin"))):
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "name": data.name, "address": data.address, "created_at": now_iso()}
    await db.outlets.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.delete("/outlets/{oid}")
async def delete_outlet(oid: str, user: dict = Depends(require_roles("umkm_admin"))):
    await db.outlets.delete_one({"id": oid, "umkm_id": user["umkm_id"]})
    return {"ok": True}


# ================= CASHIERS =================
@api_router.get("/cashiers")
async def list_cashiers(user: dict = Depends(require_roles("umkm_admin"))):
    users = await db.users.find({"umkm_id": user["umkm_id"], "role": "cashier"}).to_list(200)
    return [clean_user(u) for u in users]


@api_router.post("/cashiers")
async def create_cashier(data: CashierInput, user: dict = Depends(require_roles("umkm_admin"))):
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    res = await db.users.insert_one({"email": email, "password_hash": hash_password(data.password), "name": data.name,
                                     "role": "cashier", "umkm_id": user["umkm_id"], "outlet_id": data.outlet_id,
                                     "is_active": True, "created_at": now_iso()})
    return clean_user(await db.users.find_one({"_id": res.inserted_id}))


@api_router.delete("/cashiers/{cid}")
async def delete_cashier(cid: str, user: dict = Depends(require_roles("umkm_admin"))):
    try:
        oid = ObjectId(cid)
    except Exception:
        raise HTTPException(status_code=400, detail="ID kasir tidak valid")
    await db.users.delete_one({"_id": oid, "umkm_id": user["umkm_id"], "role": "cashier"})
    return {"ok": True}


# ================= PRODUCTS =================
@api_router.get("/products")
async def list_products(user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    return await db.products.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).sort("created_at", -1).to_list(1000)


@api_router.post("/products")
async def create_product(data: ProductInput, user: dict = Depends(require_roles("umkm_admin"))):
    doc = data.model_dump()
    doc.update({"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "created_at": now_iso()})
    await db.products.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/products/{pid}")
async def update_product(pid: str, data: ProductInput, user: dict = Depends(require_roles("umkm_admin"))):
    await db.products.update_one({"id": pid, "umkm_id": user["umkm_id"]}, {"$set": data.model_dump()})
    return await db.products.find_one({"id": pid}, {"_id": 0})


@api_router.delete("/products/{pid}")
async def delete_product(pid: str, user: dict = Depends(require_roles("umkm_admin"))):
    await db.products.delete_one({"id": pid, "umkm_id": user["umkm_id"]})
    return {"ok": True}


# ================= CUSTOMERS (kasbon) =================
@api_router.get("/customers")
async def list_customers(user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    return await db.customers.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).sort("name", 1).to_list(1000)


@api_router.post("/customers")
async def create_customer(data: CustomerInput, user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "name": data.name, "phone": data.phone,
           "note": data.note, "balance": 0, "created_at": now_iso()}
    await db.customers.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.post("/customers/{cid}/pay")
async def pay_debt(cid: str, amount: float = Query(...), user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    cust = await db.customers.find_one({"id": cid, "umkm_id": user["umkm_id"]})
    if not cust:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    new_bal = max(0, cust.get("balance", 0) - amount)
    await db.customers.update_one({"id": cid}, {"$set": {"balance": new_bal}})
    return {"ok": True, "balance": new_bal}


@api_router.delete("/customers/{cid}")
async def delete_customer(cid: str, user: dict = Depends(require_roles("umkm_admin"))):
    await db.customers.delete_one({"id": cid, "umkm_id": user["umkm_id"]})
    return {"ok": True}


# ================= SUPPLIERS =================
@api_router.get("/suppliers")
async def list_suppliers(user: dict = Depends(require_roles("umkm_admin"))):
    return await db.suppliers.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).sort("name", 1).to_list(1000)


@api_router.post("/suppliers")
async def create_supplier(data: SupplierInput, user: dict = Depends(require_roles("umkm_admin"))):
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "name": data.name, "phone": data.phone,
           "note": data.note, "created_at": now_iso()}
    await db.suppliers.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.delete("/suppliers/{sid}")
async def delete_supplier(sid: str, user: dict = Depends(require_roles("umkm_admin"))):
    await db.suppliers.delete_one({"id": sid, "umkm_id": user["umkm_id"]})
    return {"ok": True}


# ================= PURCHASES (restock) =================
@api_router.post("/purchases")
async def create_purchase(data: PurchaseInput, user: dict = Depends(require_roles("umkm_admin"))):
    if not data.items:
        raise HTTPException(status_code=400, detail="Tidak ada item pembelian")
    total = sum(i.cost * i.qty for i in data.items)
    supplier_name = None
    if data.supplier_id:
        sup = await db.suppliers.find_one({"id": data.supplier_id, "umkm_id": user["umkm_id"]})
        if sup:
            supplier_name = sup["name"]
    for i in data.items:
        # add stock and update latest cost
        await db.products.update_one({"id": i.product_id, "umkm_id": user["umkm_id"]},
                                     {"$inc": {"stock": i.qty}, "$set": {"cost": i.cost}})
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "outlet_id": data.outlet_id or user.get("outlet_id"),
           "cashier_id": user["id"], "cashier_name": user["name"], "type": "purchase",
           "items": [i.model_dump() for i in data.items], "total": total, "cogs": 0,
           "supplier_id": data.supplier_id, "supplier_name": supplier_name,
           "payment_method": data.payment_method, "status": "credit" if data.payment_method == "credit" else "paid",
           "is_credit": data.payment_method == "credit", "note": data.note, "created_at": now_iso()}
    await db.transactions.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.get("/purchases")
async def list_purchases(user: dict = Depends(require_roles("umkm_admin"))):
    return await db.transactions.find({"umkm_id": user["umkm_id"], "type": "purchase"}, {"_id": 0}).sort("created_at", -1).to_list(500)


# ================= TRANSACTIONS =================
@api_router.post("/transactions/sale")
async def create_sale(data: SaleInput, user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    if not data.items:
        raise HTTPException(status_code=400, detail="Keranjang kosong")
    subtotal = sum(i.price * i.qty for i in data.items)
    tax = round(subtotal * (data.tax_rate / 100.0), 2)
    total = round(subtotal + tax - data.discount, 2)
    cogs = sum(i.cost * i.qty for i in data.items)
    customer_name = None
    customer_phone = None
    if data.customer_id:
        cust = await db.customers.find_one({"id": data.customer_id, "umkm_id": user["umkm_id"]})
        if cust:
            customer_name = cust["name"]
            customer_phone = cust.get("phone") or None
            if data.is_credit:
                await db.customers.update_one({"id": data.customer_id}, {"$inc": {"balance": total}})
    for i in data.items:
        await db.products.update_one({"id": i.product_id, "umkm_id": user["umkm_id"]}, {"$inc": {"stock": -i.qty}})
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "outlet_id": data.outlet_id or user.get("outlet_id"),
           "cashier_id": user["id"], "cashier_name": user["name"], "type": "sale",
           "items": [i.model_dump() for i in data.items], "subtotal": subtotal, "tax_rate": data.tax_rate,
           "tax": tax, "discount": data.discount, "total": total, "cogs": cogs,
           "payment_method": data.payment_method, "status": "credit" if data.is_credit else "paid",
           "is_credit": data.is_credit, "customer_id": data.customer_id, "customer_name": customer_name,
           "customer_phone": customer_phone,
           "amount_paid": data.amount_paid, "note": data.note, "created_at": now_iso()}
    await db.transactions.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.post("/transactions/expense")
async def create_expense(data: ExpenseInput, user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    doc = {"id": str(uuid.uuid4()), "umkm_id": user["umkm_id"], "outlet_id": data.outlet_id or user.get("outlet_id"),
           "cashier_id": user["id"], "cashier_name": user["name"], "type": "expense", "items": [],
           "category": data.category, "total": data.amount, "cogs": 0, "payment_method": "cash",
           "status": "paid", "is_credit": False, "note": data.note, "created_at": now_iso()}
    await db.transactions.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.get("/transactions")
async def list_transactions(type: Optional[str] = None, limit: int = 100, user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    q = {"umkm_id": user["umkm_id"]}
    if type:
        q["type"] = type
    if user["role"] == "cashier":
        q["cashier_id"] = user["id"]
    return await db.transactions.find(q, {"_id": 0}).sort("created_at", -1).to_list(limit)


@api_router.get("/transactions/{txn_id}/receipt")
async def transaction_receipt(txn_id: str, authorization: str = Header(None), auth: str = Query(None)):
    token = authorization[7:] if authorization and authorization.startswith("Bearer ") else auth
    user = await user_from_token(token)
    if user["role"] not in ("umkm_admin", "cashier"):
        raise HTTPException(status_code=403, detail="Akses ditolak")
    txn = await db.transactions.find_one({"id": txn_id, "umkm_id": user["umkm_id"]}, {"_id": 0})
    if not txn:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    umkm = await get_umkm(user["umkm_id"])

    from reportlab.lib.pagesizes import A6
    from reportlab.lib import colors
    from reportlab.lib.units import cm, mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
    from reportlab.platypus.flowables import HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
    from PIL import Image as PILImage

    logo_path = ROOT_DIR / "assets" / "logo.png"
    store_logo_bytes = None
    slp = (umkm or {}).get("logo_image_path")
    if slp:
        try:
            store_logo_bytes, _ = get_object(slp)
        except Exception:
            store_logo_bytes = None

    def _rl_img(src, max_h_cm):
        if isinstance(src, (bytes, bytearray)):
            iw, ih = PILImage.open(io.BytesIO(src)).size
            rli = RLImage(io.BytesIO(src))
        else:
            iw, ih = PILImage.open(str(src)).size
            rli = RLImage(str(src))
        h = max_h_cm * cm
        rli.drawHeight = h
        rli.drawWidth = h * (iw / ih if ih else 3.0)
        rli.hAlign = "CENTER"
        return rli

    styles = getSampleStyleSheet()
    ctr = ParagraphStyle("ctr", parent=styles["Normal"], alignment=TA_CENTER, fontSize=8, leading=11)
    ctrb = ParagraphStyle("ctrb", parent=styles["Normal"], alignment=TA_CENTER, fontSize=12, leading=15, fontName="Helvetica-Bold")
    lft = ParagraphStyle("lft", parent=styles["Normal"], alignment=TA_LEFT, fontSize=8, leading=11)
    rgt = ParagraphStyle("rgt", parent=styles["Normal"], alignment=TA_RIGHT, fontSize=8, leading=11)
    rgtb = ParagraphStyle("rgtb", parent=styles["Normal"], alignment=TA_RIGHT, fontSize=11, leading=14, fontName="Helvetica-Bold")
    brand = ParagraphStyle("brand", parent=styles["Normal"], alignment=TA_CENTER, fontSize=7, textColor=colors.HexColor("#64748B"), leading=10)

    def money(v):
        return ("Rp {:,.0f}".format(float(v or 0))).replace(",", ".")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A6, topMargin=8, bottomMargin=8, leftMargin=10, rightMargin=10)
    E = []

    def hr():
        return HRFlowable(width="100%", thickness=0.5, color=colors.grey, dash=(2, 2))

    if store_logo_bytes:
        E.append(_rl_img(store_logo_bytes, 1.6))
        E.append(Spacer(1, 4))
    E.append(Paragraph((umkm.get("name") if umkm else None) or "Toko", ctrb))
    if umkm and umkm.get("address"):
        E.append(Paragraph(umkm["address"], ctr))
    if umkm and umkm.get("phone"):
        E.append(Paragraph("Telp: " + str(umkm["phone"]), ctr))
    E.append(Spacer(1, 4)); E.append(hr()); E.append(Spacer(1, 4))
    E.append(Paragraph(str(txn.get("created_at", ""))[:19].replace("T", " "), lft))
    E.append(Paragraph("No: " + str(txn.get("id", ""))[:8].upper(), lft))
    E.append(Paragraph("Kasir: " + (txn.get("cashier_name") or "-"), lft))
    if txn.get("customer_name"):
        E.append(Paragraph("Pelanggan: " + txn["customer_name"], lft))
    if txn.get("customer_phone"):
        E.append(Paragraph("No. HP: " + str(txn["customer_phone"]), lft))
    E.append(Spacer(1, 4)); E.append(hr()); E.append(Spacer(1, 2))
    rows = []
    for it in txn.get("items", []):
        desc = "{}<br/><font size=7 color='#666666'>{} x {}</font>".format(it.get("name", ""), it.get("qty"), money(it.get("price")))
        rows.append([Paragraph(desc, lft), Paragraph(money(float(it.get("price", 0)) * float(it.get("qty", 0))), rgt)])
    if rows:
        t = Table(rows, colWidths=[None, 26 * mm])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 1), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        E.append(t)
    E.append(Spacer(1, 2)); E.append(hr()); E.append(Spacer(1, 2))
    if txn.get("discount"):
        E.append(Paragraph("Diskon: -" + money(txn["discount"]), rgt))
    E.append(Paragraph("TOTAL: " + money(txn.get("total")), rgtb))
    E.append(Paragraph("Bayar: " + str(txn.get("payment_method", "")).upper(), rgt))
    if txn.get("payment_method") == "cash" and txn.get("amount_paid"):
        E.append(Paragraph("Tunai: " + money(txn["amount_paid"]), rgt))
        E.append(Paragraph("Kembali: " + money(max(0, float(txn["amount_paid"]) - float(txn.get("total") or 0))), rgt))
    if txn.get("is_credit"):
        E.append(Paragraph("(KASBON / Belum Lunas)", ctr))
    E.append(Spacer(1, 4)); E.append(hr()); E.append(Spacer(1, 4))
    E.append(Paragraph("Terima kasih atas kunjungan Anda", ctr))
    if logo_path.exists():
        E.append(Spacer(1, 4)); E.append(_rl_img(logo_path, 0.7))
    E.append(Paragraph("Ditenagai oleh UMKM go digital", brand))
    doc.build(E)
    buf.seek(0)
    ref = str(txn.get("id", ""))[:8]
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"inline; filename=struk-{ref}.pdf"})



# ================= DASHBOARD =================
def _match_outlet(item, outlet_id):
    if not outlet_id:
        return True
    return item.get("outlet_id") == outlet_id


@api_router.get("/dashboard/summary")
async def dashboard_summary(outlet_id: Optional[str] = None, user: dict = Depends(require_roles("umkm_admin", "cashier"))):
    all_txns = await db.transactions.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).to_list(10000)
    txns = [t for t in all_txns if _match_outlet(t, outlet_id)]
    income = sum(t["total"] for t in txns if t["type"] == "sale" and not t.get("is_credit"))
    credit_sales = sum(t["total"] for t in txns if t["type"] == "sale" and t.get("is_credit"))
    expense = sum(t["total"] for t in txns if t["type"] == "expense")
    purchases = sum(t["total"] for t in txns if t["type"] == "purchase")
    cogs = sum(t.get("cogs", 0) for t in txns if t["type"] == "sale")
    today = datetime.now(timezone.utc).date().isoformat()
    today_income = sum(t["total"] for t in txns if t["type"] == "sale" and t["created_at"][:10] == today)
    sale_count = sum(1 for t in txns if t["type"] == "sale")

    # last 7 days series
    days = [(datetime.now(timezone.utc).date() - timedelta(days=i)).isoformat() for i in range(6, -1, -1)]
    series = []
    for d in days:
        inc = sum(t["total"] for t in txns if t["type"] == "sale" and t["created_at"][:10] == d)
        exp = sum(t["total"] for t in txns if t["type"] in ("expense", "purchase") and t["created_at"][:10] == d)
        series.append({"date": d[5:], "income": inc, "expense": exp})

    # top products
    counter = {}
    for t in txns:
        if t["type"] == "sale":
            for i in t.get("items", []):
                k = i["name"]
                counter[k] = counter.get(k, 0) + i["qty"]
    top = sorted(counter.items(), key=lambda x: -x[1])[:5]
    top_products = [{"name": k, "qty": v} for k, v in top]

    all_products = await db.products.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).to_list(1000)
    products = [p for p in all_products if not outlet_id or p.get("outlet_id") in (outlet_id, None, "")]
    low_stock = [p for p in products if p.get("stock", 0) <= p.get("low_stock_threshold", 5)]
    customers = await db.customers.find({"umkm_id": user["umkm_id"]}, {"_id": 0}).to_list(1000)
    receivables = sum(c.get("balance", 0) for c in customers)

    return {"income": income, "expense": expense, "purchases": purchases, "profit": income - expense,
            "gross_profit": income - cogs, "credit_sales": credit_sales, "today_income": today_income,
            "sale_count": sale_count, "product_count": len(products), "low_stock_count": len(low_stock),
            "low_stock": low_stock[:10], "receivables": receivables, "series": series, "top_products": top_products}


# ================= REPORTS =================
def parse_range(start, end):
    return start, end


async def _report_data(umkm_id, start, end, outlet_id=None):
    q = {"umkm_id": umkm_id}
    txns = await db.transactions.find(q, {"_id": 0}).sort("created_at", 1).to_list(20000)
    if start:
        txns = [t for t in txns if t["created_at"][:10] >= start]
    if end:
        txns = [t for t in txns if t["created_at"][:10] <= end]
    if outlet_id:
        txns = [t for t in txns if t.get("outlet_id") == outlet_id]
    return txns


@api_router.get("/reports/summary")
async def report_summary(start: Optional[str] = None, end: Optional[str] = None, outlet_id: Optional[str] = None, user: dict = Depends(require_roles("umkm_admin"))):
    txns = await _report_data(user["umkm_id"], start, end, outlet_id)
    revenue = sum(t["total"] for t in txns if t["type"] == "sale")
    cogs = sum(t.get("cogs", 0) for t in txns if t["type"] == "sale")
    gross = revenue - cogs
    exp_by_cat = {}
    for t in txns:
        if t["type"] == "expense":
            exp_by_cat[t.get("category", "Lain")] = exp_by_cat.get(t.get("category", "Lain"), 0) + t["total"]
    total_expense = sum(exp_by_cat.values())
    purchases = sum(t["total"] for t in txns if t["type"] == "purchase")
    net = gross - total_expense
    cash_in = sum(t["total"] for t in txns if t["type"] == "sale" and not t.get("is_credit"))
    cash_out = total_expense + sum(t["total"] for t in txns if t["type"] == "purchase" and not t.get("is_credit"))
    return {"revenue": revenue, "cogs": cogs, "gross_profit": gross, "expenses_by_category": exp_by_cat,
            "total_expense": total_expense, "purchases": purchases, "net_profit": net, "cash_in": cash_in,
            "cash_out": cash_out, "net_cash_flow": cash_in - cash_out, "transaction_count": len(txns)}


async def user_from_token(token: str) -> dict:
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
    if not user:
        raise HTTPException(status_code=401, detail="User tidak ditemukan")
    return clean_user(user)


@api_router.get("/reports/export")
async def export_report(format: str = "excel", start: Optional[str] = None, end: Optional[str] = None,
                        outlet_id: Optional[str] = None, authorization: str = Header(None), auth: str = Query(None)):
    token = authorization[7:] if authorization and authorization.startswith("Bearer ") else auth
    user = await user_from_token(token)
    if user["role"] != "umkm_admin":
        raise HTTPException(status_code=403, detail="Akses ditolak")
    txns = await _report_data(user["umkm_id"], start, end, outlet_id)
    umkm = await get_umkm(user["umkm_id"])
    bn = umkm["name"] if umkm else "UMKM"
    type_label = {"sale": "Penjualan", "expense": "Pengeluaran", "purchase": "Pembelian"}
    rows = [{"Tanggal": t["created_at"][:19].replace("T", " "), "Tipe": type_label.get(t["type"], t["type"]),
             "Keterangan": t.get("category") or t.get("supplier_name") or ", ".join(i["name"] for i in t.get("items", [])) or "-",
             "Metode": t.get("payment_method", "-"), "Kasir": t.get("cashier_name", "-"),
             "Status": t.get("status", "-"), "Jumlah": t["total"]} for t in txns]

    logo_path = ROOT_DIR / "assets" / "logo.png"
    addr = (umkm or {}).get("address") or ""
    phone = (umkm or {}).get("phone") or ""
    generated = now_iso()[:19].replace("T", " ")
    income = sum(t["total"] for t in txns if t["type"] == "sale")
    expense = sum(t["total"] for t in txns if t["type"] == "expense")
    store_logo_bytes = None
    slp = (umkm or {}).get("logo_image_path")
    if slp:
        try:
            store_logo_bytes, _ = get_object(slp)
        except Exception:
            store_logo_bytes = None

    if format == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from PIL import Image as PILImage
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.2 * cm)
        styles = getSampleStyleSheet()
        small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748B"))

        def _rl_img(src, max_h_cm):
            if isinstance(src, (bytes, bytearray)):
                iw, ih = PILImage.open(io.BytesIO(src)).size
                rli = RLImage(io.BytesIO(src))
            else:
                iw, ih = PILImage.open(str(src)).size
                rli = RLImage(str(src))
            h = max_h_cm * cm
            rli.drawHeight = h
            rli.drawWidth = h * (iw / ih if ih else 3.0)
            return rli

        elems = []
        left = _rl_img(store_logo_bytes, 1.7) if store_logo_bytes else None
        right = _rl_img(logo_path, 1.3) if logo_path.exists() else None
        if left and right:
            htbl = Table([[left, right]], colWidths=[9 * cm, 8 * cm])
            htbl.setStyle(TableStyle([("ALIGN", (1, 0), (1, 0), "RIGHT"), ("ALIGN", (0, 0), (0, 0), "LEFT"),
                                      ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
            elems.append(htbl)
            elems.append(Spacer(1, 8))
        elif left or right:
            elems.append(left or right)
            elems.append(Spacer(1, 6))
        elems.append(Paragraph("Laporan Keuangan", styles["Title"]))
        elems.append(Paragraph(f"<b>{bn}</b>", styles["Heading3"]))
        if addr:
            elems.append(Paragraph(addr, styles["Normal"]))
        if phone:
            elems.append(Paragraph(f"Telp: {phone}", styles["Normal"]))
        elems.append(Paragraph(f"Periode: {start or 'Awal'} s/d {end or 'Sekarang'}", styles["Normal"]))
        elems.append(Paragraph(f"Dicetak: {generated}", small))
        elems.append(Spacer(1, 12))
        elems.append(Paragraph(f"Total Pemasukan: Rp {income:,.0f} | Total Pengeluaran: Rp {expense:,.0f} | Laba Bersih: Rp {income - expense:,.0f}", styles["Normal"]))
        elems.append(Spacer(1, 12))
        data = [["Tanggal", "Tipe", "Keterangan", "Metode", "Kasir", "Status", "Jumlah"]] + \
               [[r["Tanggal"], r["Tipe"], r["Keterangan"][:26], r["Metode"], r["Kasir"], r["Status"], f"Rp {r['Jumlah']:,.0f}"] for r in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2811E")),
                                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                    ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")])]))
        elems.append(table)
        elems.append(Spacer(1, 16))
        elems.append(Paragraph("Ditenagai oleh UMKM go digital - aplikasi keuangan &amp; kasir UMKM Indonesia", small))
        doc.build(elems)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename=laporan-{bn}.pdf"})
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.drawing.image import Image as XLImage
        wb = Workbook()
        ws = wb.active
        ws.title = "Laporan"
        rr = 1
        anchored = False
        if store_logo_bytes:
            try:
                from PIL import Image as PILImage
                iw, ih = PILImage.open(io.BytesIO(store_logo_bytes)).size
                simg = XLImage(io.BytesIO(store_logo_bytes))
                simg.height = 64
                simg.width = int(64 * (iw / ih)) if ih else 64
                ws.add_image(simg, "A1")
                anchored = True
            except Exception:
                pass
        if logo_path.exists():
            try:
                aimg = XLImage(str(logo_path))
                aimg.width, aimg.height = 190, 61
                ws.add_image(aimg, "F1" if anchored else "A1")
                anchored = True
            except Exception:
                pass
        if anchored:
            rr = 6
        ws.cell(row=rr, column=1, value="Laporan Keuangan").font = Font(bold=True, size=14); rr += 1
        ws.cell(row=rr, column=1, value=bn).font = Font(bold=True, size=12); rr += 1
        if addr:
            ws.cell(row=rr, column=1, value=addr); rr += 1
        if phone:
            ws.cell(row=rr, column=1, value=f"Telp: {phone}"); rr += 1
        ws.cell(row=rr, column=1, value=f"Periode: {start or 'Awal'} s/d {end or 'Sekarang'}"); rr += 1
        ws.cell(row=rr, column=1, value=f"Dicetak: {generated}"); rr += 1
        ws.cell(row=rr, column=1, value=f"Pemasukan: Rp {income:,.0f}   Pengeluaran: Rp {expense:,.0f}   Laba Bersih: Rp {income - expense:,.0f}").font = Font(bold=True); rr += 2
        headers = ["Tanggal", "Tipe", "Keterangan", "Metode", "Kasir", "Status", "Jumlah"]
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=rr, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="F2811E")
        rr += 1
        for row in rows:
            for c, h in enumerate(headers, start=1):
                ws.cell(row=rr, column=c, value=row[h])
            rr += 1
        rr += 1
        ws.cell(row=rr, column=1, value="Ditenagai oleh UMKM go digital").font = Font(italic=True, color="64748B")
        for i, w in enumerate([20, 12, 32, 10, 16, 12, 14], start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename=laporan-{bn}.xlsx"})


@api_router.get("/")
async def root():
    return {"message": "UMKM Pay API"}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.umkms.create_index("id", unique=True)
    await db.products.create_index("umkm_id")
    await db.transactions.create_index("umkm_id")
    # seed super admin
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_pw = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({"email": admin_email, "password_hash": hash_password(admin_pw), "name": "Super Admin",
                                   "role": "super_admin", "umkm_id": None, "is_active": True, "created_at": now_iso()})
    elif not verify_password(admin_pw, existing["password_hash"]) or existing.get("role") != "super_admin":
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_pw), "role": "super_admin"}})
    # seed platform settings
    if not await db.platform.find_one({"id": "platform"}):
        await db.platform.insert_one({"id": "platform", "plans": DEFAULT_PLANS, "qris_image_path": None,
                                      "business_name": "UMKM Pay", "contact": ""})
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
